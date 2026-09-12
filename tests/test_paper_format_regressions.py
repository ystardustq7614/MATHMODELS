"""Regression checks for real manuscript failures missed by S8."""
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.agents/skills/paper-formal-writer/scripts'))
from check_paper_format import internal_language_failures, char_count
from paper_scope import check_rendered_scope, delivery_scope
from format_formal_docx import configure_document, add_heading, add_table_from_rows
from validate_authoring import validate_common
from authoring_contracts import numeric_variants
from docx import Document
from docx.oxml.ns import qn


class FormatRegressionTests(unittest.TestCase):
    def test_scientific_notation_exponent_is_not_trimmed(self):
        variants = numeric_variants(1.590029219400435e-10)
        self.assertIn('1.590029219400435e-10', variants)
        self.assertNotIn('1.590029219400435e-1', variants)

    def test_three_line_table_has_no_vertical_grid(self):
        document = Document()
        configure_document(document)
        add_table_from_rows(document, [['量', '值'], ['温度', '1.2']])
        for row in document.tables[0].rows:
            for cell in row.cells:
                borders = cell._tc.find('.//' + qn('w:tcBorders'))
                self.assertEqual('nil', borders.find(qn('w:left')).get(qn('w:val')))
                self.assertEqual('nil', borders.find(qn('w:right')).get(qn('w:val')))

    def test_image_destination_is_not_rendered_prose(self):
        self.assertEqual([], internal_language_failures(
            '# 1 问题重述\n![图 1 温度](paper_output/figures/temperature.png)'))
        self.assertTrue(internal_language_failures(
            '# 1 问题重述\n读取 paper_output/results/metrics.json。'))
        self.assertTrue(internal_language_failures(
            '# 1 问题重述\n![paper_output 不应显示](figures/a.png)'))

    def test_problem_appendix_reference_does_not_end_main_paper(self):
        plan = {'delivery': delivery_scope()}
        pages = ['正文'] * 20
        pages[4] = '附录4的经验式为：\n这里仍然是模型推导。'
        failures, counts = check_rendered_scope(pages + ['附录\n代码'], plan)
        self.assertEqual([], failures)
        self.assertEqual(20, counts['counted_main_pages'])
        pages[4] = '附录4 的经验式为：\nPDF在数字后插入空格。'
        failures, counts = check_rendered_scope(pages + ['附录\n代码'], plan)
        self.assertEqual([], failures)
        self.assertEqual(20, counts['counted_main_pages'])
        for heading in ['附录 A：算法', '附录1：代码', 'Appendix A: Code']:
            _, counts = check_rendered_scope(pages + [heading], plan)
            self.assertEqual(20, counts['counted_main_pages'])

    def test_appendix_and_code_cannot_fill_body_count(self):
        body = '# 摘要\n正文内容'
        self.assertEqual(char_count(body), char_count(body + '\n# 附录\n' + '代码' * 1000))
        self.assertEqual(char_count(body), char_count(body + '\n```python\n' + 'a=1\n' * 1000 + '```'))

    def test_a4_page_numbers_and_abstract_boundary(self):
        document = Document()
        configure_document(document)
        page = document.sections[0]
        self.assertAlmostEqual(21, page.page_width.cm, places=2)
        self.assertAlmostEqual(29.7, page.page_height.cm, places=2)
        self.assertTrue(page.footer._element.findall('.//' + qn('w:fldSimple')))
        add_heading(document, '1 问题重述', 1)
        self.assertTrue(document.paragraphs[-1].paragraph_format.page_break_before)

    def test_appendix_code_is_not_prose_or_numeric_evidence(self):
        spec = {'target_chars': 1, 'requirements': {'evidence': [
            {'evidence_id': 'metric:test', 'value': 2468, 'unit': 'kg'}]}}
        issues, _ = validate_common('# 标题\n正文\n```python\n# TODO\nx = 2468 # kg\n```',
                                    spec, 'FINAL', require_marker=False)
        self.assertEqual(['numeric-evidence'], [i['category'] for i in issues])

    def test_table_numbers_do_not_inherit_body_indent(self):
        document = Document()
        configure_document(document)
        add_table_from_rows(document, [['含水率'], ['1.7172']])
        cell = document.tables[0].cell(1, 0)
        self.assertEqual(0, cell.paragraphs[0].paragraph_format.first_line_indent)
        self.assertIsNotNone(cell._tc.find('.//' + qn('w:noWrap')))


if __name__ == '__main__':
    unittest.main()
